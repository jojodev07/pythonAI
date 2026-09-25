import os
import io
from urllib.parse import quote
from typing import List
import asyncio
from fastapi import FastAPI, HTTPException, status
from fastapi.responses import StreamingResponse
from fastapi.middleware.cors import CORSMiddleware
from fastapi import BackgroundTasks
from openai import AsyncOpenAI
import chromadb
from fastapi.middleware.cors import CORSMiddleware
from schemas import QueryRequest, QueryResponse, StageDetail, LessonPlanStages, LessonPlanSchema, LessonPlanRequest
from dotenv import load_dotenv

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.enum.table import WD_TABLE_DIRECTION
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

origins = [
    "http://localhost:5173",
    "http://localhost:8080",
    "https://springbackend-zei7.onrender.com",
    "https://teachassist-delta.vercel.app"
]

app = FastAPI(title="RAG API with gpt 4o mini")
app.add_middleware(
    CORSMiddleware,
    allow_origins=origins,
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

load_dotenv("minimal.env")
# Initialize clients (Ensure OPENAI_API_KEY is set in .env):
openai_client = AsyncOpenAI(api_key=os.getenv("OPENAI_API_KEY"))

# Connect to vector db:

try:
    clientdb = chromadb.CloudClient(
    api_key=os.getenv("CHROMADB_KEY"),
    tenant=os.getenv("CHROMADB_TENANT"),
    database='prod'
    )
    collection = clientdb.get_collection("arabic_books")
except Exception as e:
    print(f"Failed to connect to Chroma Cloud: {str(e)}")
    collection = None

async def retrieve_context(query: str, top_k: int) -> list[str]:
    if not collection:
        return []
    try:
        embedding_response = await openai_client.embeddings.create(
            model="text-embedding-3-small",
            input=query
        )
        query_vector = embedding_response.data[0].embedding

        results = await asyncio.to_thread(
            collection.query,
            query_embeddings=[query_vector],
            n_results=top_k,
            include=["documents", "metadatas"]
        )
        documents = results.get("documents", [[]])

        return documents[0] if documents else []
    except Exception as e:
        print(f"Error retrieving context: {str(e)}")
        return []


async def generate_answer(question: str, contexts: list[str]) -> str:
    source_knowledge = "\n---\n".join(contexts)

    # 2. Updated system prompt to include source_knowledge
    system_prompt = (
            "أنت المستشار الرقمي للمعلم، نظام خبير ذكي ومبادرة تربوية تهدف إلى مساندة "
            "المعلمين والمعلمات في المدارس لحل مشكلاتهم اليومية (التربوية، السلوكية، والقانونية الإدارية) "
            "فوراً، وبما يتوافق تماماً مع رؤية ورسالة وقوانين وزارة التربية والتعليم الرسمية المرفقة في ملفات المعرفة لديك.\n\n"

            f"السياق المتاح:\n{source_knowledge}\n\n"

            "قواعد الإجابة:\n"
            "1. أجب باللغة العربية الواضحة والمهنية.\n"
            "2. استخدم Markdown القياسي فقط لتنسيق الإجابة، ولا تستخدم HTML.\n"
            "3. ابدأ الإجابة بعنوان Markdown قصير باستخدام # أو ## عند الحاجة.\n"
            "4. بعد العنوان، أضف ملخصاً موجزاً من جملة أو جملتين يوضح الإجابة.\n"
            "5. استخدم القوائم النقطية (-) أو (*) عند عرض عدة نقاط.\n"
            "6. استخدم القوائم المرقمة فقط عندما يكون ترتيب الخطوات مهماً.\n"
            "7. استخدم **النص العريض** لتوضيح المصطلحات أو النقاط المهمة.\n"
            "8. استخدم `inline code` فقط عند الحاجة لعرض أسماء حقول أو أكواد تقنية.\n"
            "9. استخدم الجداول Markdown فقط عندما تكون المقارنة أو البيانات الجدولية مفيدة.\n"
            "10. لا تضع كل الإجابة داخل كتلة كود ```.\n"
            "11. لا تستخدم HTML مثل <div> أو <p> أو <br> أو inline styles.\n"
            "12. لا تضف Markdown غير ضروري؛ اجعل التنسيق واضحاً ونظيفاً وسهل القراءة.\n\n"

            "قواعد المحتوى:\n"
            "إذا كانت المعلومات المتاحة غير كافية للإجابة، فاعتذر للمستخدم وأوضح أن البيانات المتاحة "
            "غير كافية للإجابة بدقة، ولا تخترع أو تفترض معلومات غير موجودة.\n"
            "إذا كان السؤال لا يتعلق بالتعليم أو علم النفس أو التربية، فاعتذر للمستخدم وأوضح أن "
            "المحادثة مخصصة لهذه المجالات فقط.\n\n"

            "قواعد التوثيق (الاستشهاد بالمصادر):\n"
            "13. عند الاستناد إلى معلومة محددة من 'السياق المتاح'، حاول استخراج اسم المؤلف أو "
            "الجهة المُصدرة (مثل اسم الوزارة، أو اسم الباحث، أو عنوان الوثيقة الرسمية) وسنة النشر إن وُجدت "
            "ضمن النص المرفق.\n"
            "14. إذا توفر اسم المؤلف والسنة، استخدم أسلوب التوثيق (APA) بصيغته العربية داخل النص، "
            "مثل: (اسم المؤلف، السنة) أو حسب اسم المؤلف (السنة)، مثال: (وزارة التربية والتعليم، 2023).\n"
            "15. إذا توفر اسم المؤلف فقط دون سنة، اذكر اسم المؤلف فقط دون اختراع سنة، مثال: "
            "(وزارة التربية والتعليم).\n"
            "16. لا تخترع أو تفترض اسم مؤلف أو سنة غير مذكورة صراحة في السياق المتاح؛ إذا لم تتوفر "
            "أي معلومة توثيقية، اذكر المعلومة دون استشهاد بدلاً من تأليف مصدر غير حقيقي.\n"
        )

    try:
        response = await openai_client.chat.completions.create(
            model="gpt-4o-mini",
            messages=[
                {"role": "system", "content": system_prompt},
                {"role": "user", "content": question},
            ],
            temperature=0.1
        )

        # 3. Fixed response indexing [0]
        return response.choices[0].message.content
    except Exception as e:
        print(f"Error generating answer: {str(e)}")
        # 4. Return string fallback instead of None to keep Pydantic happy
        return "حدث خطأ أثناء توليد الإجابة. يرجى المحاولة لاحقاً."
    
    
async def generate_lesson_plan_service(req: LessonPlanRequest) -> LessonPlanSchema:

    system_prompt = """
    أنت مساعد تربوي خبير في المناهج الأردنية.
    مهمتك هي إعداد خطة درس نموذجية مطابقة تماماً لدليل أداة تخطيط الدروس الصادر عن وزارة التربية والتعليم الأردنية.
    
    شروط التوليد:
    1. التعلم القبلي: حدد المعارف السابقة والخبرات المسبقة ورابطها بالدرس.
    2. النتاجات التعليمية: صيغ نتاجات SMART تبدأ بأفعال سلوكية (يُعرّف، يُحلّل، يميّز، يثمن).
    3. دور المعلم والمتعلم: حدد المهام بوضوح لكل مرحلة.
    4. مراعاة مستويات الطلبة (دون المتوسط، المتوسط، فوق المتوسط) في مرحلة الشرح والتفصيل.
    5. التوسع ودعم التميز: أنشطة تفكير عليا ودعم الأقران.
    6. اللغة: لغة عربية سليمة ورسمية.
    """

    user_prompt = (
    f"المبحث: {req.subject} | "
    f"عنوان الوحدة: {req.unit_title} | "
    f"موضوع الدرس: {req.lesson_title} | "
    f"الصف: {req.grade} | "
    f"عدد الحصص: {req.periods}"
)

    # Use client.beta.chat.completions.parse for guaranteed structure
    completion = await openai_client.beta.chat.completions.parse(
        model="gpt-4o-mini",
        messages=[
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": user_prompt}
        ],
        response_format=LessonPlanSchema,
        temperature=0.3
    )

    return completion.choices[0].message.parsed

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Sets inner cell padding in dxa (1 pt = 20 dxa)."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for margin_name, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{margin_name}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_background(cell, hex_color: str):
    """Applies background shading to a table cell."""
    shading_xml = f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>'
    cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))

def build_docx_in_memory(data: LessonPlanSchema) -> io.BytesIO:
    doc = Document()

    # Configure 0.5 inch document margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.5)
        section.right_margin = Inches(0.5)

    # Document Header
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run("خطة الدرس")
    run_title.font.name = "Traditional Arabic"
    run_title.font.size = Pt(18)
    run_title.font.bold = True

    # Header Top Metadata Table (RTL)
    top_table = doc.add_table(rows=2, cols=3)
    top_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    top_table.table_direction = WD_TABLE_DIRECTION.RTL
    
    # Fill Header Text
    row0 = top_table.rows[0].cells
    row0[0].text = f"المبحث: {data.subject}"
    row0[1].text = f"عنوان الوحدة: {data.unit_title}"
    row0[2].text = f"موضوع الدرس: {data.lesson_title}"

    row1 = top_table.rows[1].cells
    row1[0].text = f"عدد الحصص: {data.lessons_count}"
    row1[1].text = f"التعلم القبلي: {data.prior_learning}"
    row1[2].text = "النتاجات التعليمية:\n" + "\n".join([f"- {o}" for o in data.learning_outcomes])

    # Right Side Grid Table (Attendance, Section, Date)
    grid_table = doc.add_table(rows=4, cols=9)
    grid_table.alignment = WD_TABLE_ALIGNMENT.RIGHT
    grid_table.table_direction = WD_TABLE_DIRECTION.RTL
    grid_table.style = 'Table Grid'

    labels = ["الصف/الشعبة:", "عدد الغياب/العدد الكلي", "ترتيب الحصة/الحصص", "اليوم والتاريخ:"]
    for i, label in enumerate(labels):
        cell = grid_table.rows[i].cells[8]
        cell.text = label

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Main Stages Table
    main_table = doc.add_table(rows=5, cols=4)
    main_table.style = 'Table Grid'
    main_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    main_table.table_direction = WD_TABLE_DIRECTION.RTL

    # Column Titles
    headers = ["مراحل الحصة", "دور المعلم Teacher Action", "دور المتعلم Learner Action", "الزمن للمرحلة"]
    hdr_cells = main_table.rows[0].cells
    for i, text in enumerate(headers):
        hdr_cells[i].text = text
        set_cell_background(hdr_cells[i], "F0F0F0")
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.runs[0].font.bold = True
        hdr_cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    # Table Stage Rows
    stages_data = [
        ("1- التهيئة والاندماج\nEngagement", data.stages.engagement),
        ("2- الشرح والتفصيل مع مراعاة مستويات الطلبة (التعليم المتمايز)\nExplanation", data.stages.explanation),
        ("3- التوسع ودعم التميز\nElaboration", data.stages.elaboration),
        ("تأكيد التعلم (الغلق)\nClosing", data.stages.closing),
    ]

    for row_idx, (stage_title, stage_obj) in enumerate(stages_data, start=1):
        cells = main_table.rows[row_idx].cells
        cells[0].text = stage_title
        cells[1].text = stage_obj.teacher_action
        cells[2].text = stage_obj.learner_action
        cells[3].text = str(stage_obj.time_minutes)

        # Center-align stage names and time columns
        cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cells[3].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        for c in cells:
            set_cell_margins(c, top=140, bottom=140, left=140, right=140)
            c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    # Set Column Width Ratios
    col_widths = [Inches(2.0), Inches(2.8), Inches(2.8), Inches(0.8)]
    for row in main_table.rows:
        for idx, width in enumerate(col_widths):
            row.cells[idx].width = width

    # Self-Reflection Footer Section
    doc.add_paragraph().paragraph_format.space_before = Pt(12)
    
    ref_table = doc.add_table(rows=1, cols=1)
    ref_table.style = 'Table Grid'
    ref_table.table_direction = WD_TABLE_DIRECTION.RTL
    ref_cell = ref_table.rows[0].cells[0]
    ref_cell.text = f"التأمل الذاتي : حول عمليتي التعلم والتعليم Self-Reflection on Learning and Teaching\n{data.self_reflection}"
    set_cell_margins(ref_cell, top=140, bottom=140, left=140, right=140)

    # Signatures
    doc.add_paragraph().paragraph_format.space_before = Pt(12)
    sig_p = doc.add_paragraph("اسم المعلم وتوقيعه :                                              توقيع المشرف التربوي :                                             مدير المدرسة: ")
    sig_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    target_stream = io.BytesIO()
    doc.save(target_stream)
    target_stream.seek(0)
    return target_stream


def cleanup_stream(stream: io.BytesIO):
    """إغلاق وتفريغ ذاكرة الـ stream بعد إتمام التحميل."""
    stream.close()

@app.post("/ai/v1/query", response_model=QueryResponse)
async def rag_endpoint(payload: QueryRequest):
    contexts = await retrieve_context(payload.question, payload.top_k)

    if contexts == []:
        return QueryResponse(answer="غير متواجد في قاعدة البيانات. يرجى المحاولة بسؤال آخر" , context=[])
    answer = await generate_answer(payload.question, contexts)
    
    return QueryResponse(answer=answer, context=contexts)

@app.post("/api/generate-plan", response_model=LessonPlanSchema)
async def generate_plan_endpoint(payload: LessonPlanRequest):
    """Generates JSON payload for front-end rendering or previewing."""
    return await generate_lesson_plan_service(payload)

@app.post("/api/download-docx")
async def download_docx_endpoint(payload: LessonPlanRequest, background_tasks: BackgroundTasks):
    """Generates lesson plan and returns downloadable .docx file stream."""
    plan_data = await generate_lesson_plan_service(payload)
    
    # Run sync document creation inside a thread pool
    file_stream = await asyncio.to_thread(build_docx_in_memory, plan_data)
    
    background_tasks.add_task(cleanup_stream, file_stream)
    
    filename = f"lesson_plan_{payload.lesson_title}.docx"
    encoded_filename = quote(filename, safe="")
    return StreamingResponse(
        file_stream,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={
            "Content-Disposition": (
                f'attachment; filename="lesson_plan.docx"; '
                f"filename*=UTF-8''{encoded_filename}"
            )
        }
    )
# keep service alive:
@app.get("/", response_model=str, status_code=status.HTTP_200_OK)
@app.head("/", response_model=str, status_code=status.HTTP_200_OK)
async def sayHi():
    return "Hello! I'm alive :p"